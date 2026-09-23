import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = docx.Document()

# Page Setup
for section in doc.sections:
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

# Title
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_title = p_title.add_run("DARPA TECHNICAL PROGRAM PROPOSAL\nPROJECT AEGIS-MONAD")
run_title.bold = True
run_title.font.size = Pt(18)
run_title.font.color.rgb = RGBColor(15, 23, 42)

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_sub = p_sub.add_run("Autonomous Enclave Formal Verification & Grid-Isolated SOC Infrastructure\n")
run_sub.italic = True
run_sub.font.size = Pt(11)
run_sub.font.color.rgb = RGBColor(71, 85, 105)

# Meta Table
table = doc.add_table(rows=4, cols=1)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
meta_info = [
    ("Proposed Program Manager:", " Asaad Morman (D.A.S. Candidate, Bowie State University)"),
    ("Target Offices:", " DARPA Information Innovation Office (I2O) / Defense Sciences Office (DSO)"),
    ("Classification:", " UNCLASSIFIED // CUI // Defense Research Baseline"),
    ("Digital Provenance Invariant:", " HW_KEY_0x889_BSU_DAS_2026_EDR")
]
for i, (label, val) in enumerate(meta_info):
    cell = table.cell(i, 0)
    p = cell.paragraphs[0]
    r1 = p.add_run(label)
    r1.bold = True
    r2 = p.add_run(val)

doc.add_paragraph().paragraph_format.space_after = Pt(12)

# Heilmeier Content
questions = [
    ("1. What are you trying to do? (No Jargon Statement)", 
     "We are building a defense-grade Security Operations Center (SOC) artificial intelligence that cannot be tricked into leaking classified secrets or executing unauthorized actions. The system mathematically screens every single thought and decision before the AI outputs it, operates inside encrypted hardware vaults, and manages its own off-grid power and heat recovery for tactical deployment in hostile environments."),
    
    ("2. How is it done today, and what are the limits of current practice?",
     "• Probabilistic Guardrails: Current AI safety in SOCs relies on post-hoc filtering, RLHF, or system instructions. These methods are soft and probabilistic—adversarial prompt injections reliably bypass them, leading to a non-zero probability of CUI leakage (P_violation > 0).\n"
     "• Separated Hardware Isolation: Existing architectures treat software TEEs and vector databases as disconnected layers, leaving embeddings vulnerable to inversion attacks.\n"
     "• Grid Dependency: High-performance AI hardware clusters require fixed commercial power grids with no integrated thermal or microgrid adaptation for expeditionary edge posts."),
    
    ("3. What is new in your approach and why do you think it will be successful?",
     "AEGIS-MONAD shifts AI safety from probabilistic filtering to deterministic formal mathematical proof integrated directly into token decoding:\n\n"
     "• Monad SMT Logit Transformation Operator: Before softmax sampling, candidate logits L_i are constrained deterministically by a Z3 SMT logic solver graph:\n"
     "     L_hat_i = L_i + log Phi(v_i)\n"
     "Where Phi(v_i) evaluates to 1 if SMT_Verify(v_i) = SAT, and 0 if UNSAT. If a token violates security policy, log(0) = -infinity, reducing selection probability to exactly zero (P_violation = 0).\n\n"
     "• Steganographic Provenance & Differential Privacy: Vectors are perturbed via Gaussian noise bounded by (epsilon, delta)-DP and embedded with zero-width steganographic markers (U+200B/U+200C) bound to HW_KEY_0x889_BSU_DAS_2026_EDR.\n\n"
     "• Datacenter & Microgrid Telemetry Twin: Real-time optimization balances compute draw against off-grid solar, battery reserves, and direct immersion heat capture using the Grid Isolation Index (GI)."),
    
    ("4. Who cares? If you are successful, what difference will it make?",
     "• DoD & Intelligence Community: Enables immediate, safe deployment of autonomous LLM agents in CUI/Classified environments without fear of prompt injection or data spillage.\n"
     "• Expeditionary Command Posts: Gives forward units an AI SOC command center capable of operating 90%+ off-grid while capturing compute exhaust heat to support local microgrids.\n"
     "• Defense Industrial Base (DIB): Establishes the nation's first hardware-attested, formally verified standard for AI governance."),
    
    ("5. What are the risks and the payoffs?",
     "• Payoff: The DoD acquires a provably secure, zero-violation AI reasoning and infrastructure framework—moving from reactive patch security to mathematical immunity.\n"
     "• Risk & Mitigation (Latency): Z3 SMT evaluation adds latency -> Enforce a 5ms thread-isolated timeout and parallelize SMT context checks across multi-threaded C++ workers.\n"
     "• Risk & Mitigation (Memory Faults): C++ native memory faults under concurrency -> Demonstrated thread-local z3.Context() isolation, eliminating cross-thread race conditions.")
]

for title, text in questions:
    h = doc.add_paragraph()
    rh = h.add_run(title)
    rh.bold = True
    rh.font.size = Pt(12)
    rh.font.color.rgb = RGBColor(30, 58, 138)
    
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(8)

# Save Word Doc
doc.save("DARPA_AEGIS_MONAD_Proposal.docx")
print("[SUCCESS] Word Document generated successfully: DARPA_AEGIS_MONAD_Proposal.docx")
